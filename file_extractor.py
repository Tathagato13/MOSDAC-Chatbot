import os
import fitz  # For PDFs
import docx  # For .docx files
import openpyxl  # For .xlsx files

def extract_text_from_file(file_path):
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".xlsx":
        return extract_text_from_xlsx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

def extract_text_from_pdf(pdf_path):
    doc = fitz.open(pdf_path)
    text = ""
    for page_num, page in enumerate(doc):
        page_text = page.get_text()
        text += f"\n--- Page {page_num + 1} ---\n{page_text}"
    return text

def extract_text_from_docx(docx_path):
    doc = docx.Document(docx_path)
    return "\n".join([para.text for para in doc.paragraphs if para.text.strip() != ""])

def extract_text_from_xlsx(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    text = ""
    for sheet in wb.worksheets:
        text += f"\n--- Sheet: {sheet.title} ---\n"
        for row in sheet.iter_rows(values_only=True):
            line = "\t".join([str(cell) if cell is not None else "" for cell in row])
            text += line + "\n"
    return text


# Loop through the MAIN folder and ALL its sub-folders and process each file

root_folder = r"C:\Users\tatha\OneDrive\Desktop\chatbot sovan1\chatbot/files"

for dirpath, dirnames, filenames in os.walk(root_folder):
    for filename in filenames:
        file_path = os.path.join(dirpath, filename)

        try:
            text = extract_text_from_file(file_path)
            print(f"\n================= {file_path} =================")
            print(text)
            print("\n" + "="*90 + "\n")
        except Exception as e:
            print(f"Skipping {file_path} -> {e}")
